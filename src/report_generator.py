"""Generador del Reporte Ejecutivo Word (.docx) de EPS Ilo S.A.

Consume los contratos de las fases previas:

* ``economic_summary`` de :func:`src.economic_evaluator.evaluate_economic_impact`.
* ``anomalies_summary`` de
  :func:`src.anomaly_detector.summarize_anomalies_by_sector`.
* ``roi_simulation`` de
  :func:`src.economic_evaluator.simulate_anf_reduction_roi`.

El documento resultante se construye integramente en memoria y se retorna
como ``bytes`` para alimentar directamente ``st.download_button`` sin tocar
el disco del servidor.
"""

from __future__ import annotations

import io
import math
from collections.abc import Sequence
from datetime import datetime
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .config import NORMATIVE_LIMITS, TARIFF_STRUCTURE

EPS_NAME = "EPS Ilo S.A."
REPORT_TITLE = "Reporte Ejecutivo de Agua No Facturada (ANF)"
REPORT_SUBTITLE = (
    "Sistema Integrado de Operaciones, Control de ANF e Inteligencia "
    "Comercial - Ilo, Moquegua"
)

# Paleta de impresion: textos oscuros sobre fondo blanco para maximo contraste.
_NAVY = RGBColor(0x0F, 0x17, 0x2A)
_CYAN_DARK = RGBColor(0x0E, 0x74, 0x90)
_GRAY = RGBColor(0x47, 0x55, 0x69)
_RED_DARK = RGBColor(0xB9, 0x1C, 0x1C)
_GREEN_DARK = RGBColor(0x15, 0x80, 0x3D)
_AMBER_DARK = RGBColor(0x92, 0x40, 0x0E)

_SEMAFORO_FILL = {
    "VERDE": "15803D",
    "AMARILLO": "B45309",
    "ROJO": "B91C1C",
}


# ---------------------------------------------------------------------------
# Utilidades de formato y coercion
# ---------------------------------------------------------------------------
def _coerce_float(value: Any, default: float = 0.0) -> float:
    """Convierte un valor a float finito, con valor de respaldo."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return default
    if not math.isfinite(number):
        return default
    return number


def _fmt_number(value: Any, decimals: int = 0, unit: str = "") -> str:
    """Formatea un numero con separador de miles y unidad opcional."""
    number = _coerce_float(value)
    formatted = f"{number:,.{decimals}f}".replace(",", " ")
    return f"{formatted}{unit}"


def _fmt_money(value: Any, decimals: int = 0) -> str:
    """Formatea un monto en soles peruanos."""
    return f"S/ {_fmt_number(value, decimals=decimals)}"


def _fmt_pct(value: Any, decimals: int = 1) -> str:
    """Formatea un porcentaje."""
    return f"{_coerce_float(value):.{decimals}f}%"


def _fmt_payback(value: Any) -> str:
    """Formatea el payback en meses, contemplando el caso infinito."""
    number = _coerce_float(value, default=float("inf"))
    if not math.isfinite(number):
        return "No recuperable"
    return f"{number:.1f} meses"


def _fmt_int(value: Any) -> str:
    """Formatea un entero con separador de miles."""
    return _fmt_number(int(round(_coerce_float(value))), decimals=0)


# ---------------------------------------------------------------------------
# Utilidades de python-docx
# ---------------------------------------------------------------------------
def _set_cell_shading(cell, hex_color: str) -> None:
    """Aplica color de fondo a una celda de tabla."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    tc_pr.append(shading)


def _set_cell_margins(
    cell,
    top: int = 55,
    start: int = 85,
    bottom: int = 55,
    end: int = 85,
) -> None:
    """Ajusta los margenes internos de una celda (unidad: twips)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for tag, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _mark_header_row(row) -> None:
    """Marca la fila como encabezado repetible entre paginas."""
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_table_fixed_layout(table) -> None:
    """Fuerza un layout de tabla fijo para respetar los anchos definidos."""
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    """Agrega un titulo de seccion con linea inferior y color institucional."""
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.keep_with_next = True
    for run in heading.runs:
        run.font.color.rgb = _NAVY if level == 1 else _CYAN_DARK
        run.font.name = "Calibri"
        run.font.bold = True
    if level == 1:
        paragraph_format = heading.paragraph_format
        paragraph_format.space_before = Pt(16)
        paragraph_format.space_after = Pt(6)
        p_pr = heading._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), "0E7490")
        borders.append(bottom)
        p_pr.append(borders)
    else:
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(4)


def _add_paragraph(
    document: Document,
    text: str,
    *,
    size: float = 10.5,
    bold: bool = False,
    color: RGBColor | None = None,
    align: WD_ALIGN_PARAGRAPH | None = None,
    space_after: float = 6.0,
) -> None:
    """Agrega un parrafo con estilo consistente."""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.12


def _add_bullet(document: Document, text: str) -> None:
    """Agrega una vineta compatible con la plantilla por defecto."""
    paragraph = document.add_paragraph(text, style="List Bullet")
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = _GRAY
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.1


def _add_table(
    document: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    *,
    widths_cm: Sequence[float] | None = None,
    font_size: float = 8.3,
    header_fill: str = "0F172A",
    zebra: bool = True,
    align_right_from: int | None = 1,
    total_row: bool = False,
) -> None:
    """Agrega una tabla profesional con encabezado oscuro y filas zebra."""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_fixed_layout(table)

    if widths_cm:
        for index, width in enumerate(widths_cm):
            table.columns[index].width = Cm(width)

    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(str(header))
        run.font.name = "Calibri"
        run.font.size = Pt(font_size)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        _set_cell_shading(cell, header_fill)
        _set_cell_margins(cell)
        if widths_cm:
            cell.width = Cm(widths_cm[index])
    _mark_header_row(table.rows[0])

    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        is_total = total_row and row_index == len(rows) - 1
        for col_index, value in enumerate(row_values):
            cell = cells[col_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(str(value))
            run.font.name = "Calibri"
            run.font.size = Pt(font_size)
            run.font.bold = is_total
            run.font.color.rgb = _NAVY if is_total else _GRAY
            if (
                align_right_from is not None
                and col_index >= align_right_from
                and col_index > 0
            ):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            _set_cell_margins(cell)
            if widths_cm:
                cell.width = Cm(widths_cm[col_index])
            if is_total:
                _set_cell_shading(cell, "E2E8F0")
            elif zebra and row_index % 2 == 1:
                _set_cell_shading(cell, "F1F5F9")

    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_semaforo_banner(
    document: Document, semaforo: str, anf_pct: float
) -> None:
    """Inserta la banda visual del semaforo normativo SUNASS."""
    semaforo = str(semaforo or "SIN DATO").upper()
    fill = _SEMAFORO_FILL.get(semaforo, "334155")

    if semaforo == "ROJO":
        detail = (
            f"ANF de {anf_pct:.1f}% supera el umbral critico de "
            f"{NORMATIVE_LIMITS['anf_alerta_critica_pct']:.1f}%. "
            "Se requiere intervencion inmediata."
        )
    elif semaforo == "AMARILLO":
        detail = (
            f"ANF de {anf_pct:.1f}% supera la meta de "
            f"{NORMATIVE_LIMITS['anf_alerta_amarilla_pct']:.1f}%. "
            "Se recomienda un plan de control intensivo."
        )
    else:
        detail = (
            f"ANF de {anf_pct:.1f}% se mantiene dentro de la meta referencial "
            f"de {NORMATIVE_LIMITS['anf_alerta_amarilla_pct']:.1f}%."
        )

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(f"SEMÁFORO SUNASS: {semaforo}")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)

    detail_paragraph = cell.add_paragraph()
    detail_run = detail_paragraph.add_run(detail)
    detail_run.font.name = "Calibri"
    detail_run.font.size = Pt(9.5)
    detail_run.font.color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
    detail_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail_paragraph.paragraph_format.space_after = Pt(2)
    _set_cell_shading(cell, fill)
    _set_cell_margins(cell, top=120, bottom=120, start=120, end=120)

    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_page_number(paragraph) -> None:
    """Inserta un campo PAGE dinamico en el pie de pagina."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = _GRAY


def _configure_styles(document: Document) -> None:
    """Configura tipografia, margenes y estilos base del documento."""
    for section in document.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _GRAY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _NAVY


def _configure_header_footer(document: Document) -> None:
    """Agrega encabezado institucional y pie con numeracion."""
    section = document.sections[0]

    header_paragraph = section.header.paragraphs[0]
    header_run = header_paragraph.add_run(
        f"{EPS_NAME} · SIOC · Gerencia de Operaciones"
    )
    header_run.font.name = "Calibri"
    header_run.font.size = Pt(8.5)
    header_run.font.bold = True
    header_run.font.color.rgb = _CYAN_DARK
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run(
        "Documento generado automaticamente por el SIOC de EPS Ilo S.A. · Página "
    )
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = _GRAY
    _add_page_number(footer_paragraph)


# ---------------------------------------------------------------------------
# Preparacion de datos del reporte
# ---------------------------------------------------------------------------
def _sector_frame(economic_summary: dict) -> pd.DataFrame:
    """Extrae y normaliza el desagregado sectorial del resumen economico."""
    breakdown = economic_summary.get("sector_breakdown")
    if not isinstance(breakdown, pd.DataFrame) or breakdown.empty:
        return pd.DataFrame()
    frame = breakdown.copy()
    for column in (
        "volumen_producido_m3",
        "volumen_facturado_m3",
        "anf_m3",
        "anf_pct",
        "perdida_fisica_m3",
        "perdida_comercial_m3",
        "perdida_economica_soles",
        "costo_energia_quimicos_soles",
    ):
        if column in frame.columns:
            frame[column] = pd.to_numeric(frame[column], errors="coerce").fillna(0.0)
    return frame


def _anomaly_frame(anomalies_summary: Any) -> pd.DataFrame:
    """Normaliza el resumen de anomalias a un DataFrame seguro."""
    if isinstance(anomalies_summary, pd.DataFrame):
        frame = anomalies_summary.copy()
    elif anomalies_summary is None:
        return pd.DataFrame()
    else:
        frame = pd.DataFrame(anomalies_summary)
    if frame.empty:
        return frame
    for column in (
        "registros",
        "alertas_totales",
        "alertas_porcentaje",
        "fugas_fisicas",
        "fraudes_comerciales",
        "alertas_cloro",
        "alertas_turbiedad",
        "alertas_presion",
        "alertas_estadisticas",
        "anf_promedio_pct",
        "score_promedio",
        "score_maximo",
    ):
        if column in frame.columns:
            frame[column] = pd.to_numeric(frame[column], errors="coerce").fillna(0.0)
    return frame


def _semaforo_explanation(semaforo: str) -> str:
    """Explica el significado operativo del semaforo."""
    if semaforo == "ROJO":
        return (
            "Nivel critico: la perdida supera el 35% del volumen producido. "
            "Priorizar distritos con mayor perdida economica, intervenir fugas "
            "nocturnas y auditar la macromedicion."
        )
    if semaforo == "AMARILLO":
        return (
            "Nivel de alerta: el ANF supera la meta del 28%. Ejecutar el plan "
            "de reduccion, reforzar el catastro comercial y sectorizar el "
            "control de fugas."
        )
    return (
        "Nivel eficiente: el ANF se mantiene por debajo de la meta del 28%. "
        "Sostener el monitoreo y consolidar las acciones de control."
    )


# ---------------------------------------------------------------------------
# API publica
# ---------------------------------------------------------------------------
def generate_executive_docx_report(
    economic_summary: dict,
    anomalies_summary: pd.DataFrame,
    roi_simulation: dict,
) -> bytes:
    """Construye el Reporte Ejecutivo ANF en formato Word en memoria.

    Parameters
    ----------
    economic_summary:
        Diccionario retornado por
        :func:`src.economic_evaluator.evaluate_economic_impact`.
    anomalies_summary:
        DataFrame retornado por
        :func:`src.anomaly_detector.summarize_anomalies_by_sector`.
    roi_simulation:
        Diccionario retornado por
        :func:`src.economic_evaluator.simulate_anf_reduction_roi`.

    Returns
    -------
    bytes
        Contenido binario del archivo ``.docx`` listo para descarga.
    """
    if not isinstance(economic_summary, dict):
        raise TypeError("economic_summary debe ser un diccionario.")
    if not isinstance(roi_simulation, dict):
        raise TypeError("roi_simulation debe ser un diccionario.")

    sector_frame = _sector_frame(economic_summary)
    anomaly_frame = _anomaly_frame(anomalies_summary)

    producido = _coerce_float(economic_summary.get("volumen_producido_total_m3"))
    facturado = _coerce_float(economic_summary.get("volumen_facturado_total_m3"))
    anf_total = _coerce_float(economic_summary.get("anf_total_m3"))
    anf_pct = _coerce_float(economic_summary.get("anf_porcentaje"))
    perdida_fisica = _coerce_float(economic_summary.get("perdida_fisica_m3"))
    perdida_comercial = _coerce_float(economic_summary.get("perdida_comercial_m3"))
    perdida_economica = _coerce_float(
        economic_summary.get("perdida_economica_soles")
    )
    costo_energia = _coerce_float(
        economic_summary.get("costo_energia_quimicos_perdido_soles")
    )
    semaforo = str(economic_summary.get("semaforo_sunass") or "SIN DATO").upper()

    roi_recuperado = _coerce_float(roi_simulation.get("m3_recuperados_mes"))
    roi_ahorro_mensual = _coerce_float(roi_simulation.get("ahorro_mensual_soles"))
    roi_ahorro_anual = _coerce_float(roi_simulation.get("ahorro_anual_soles"))
    roi_payback = roi_simulation.get("payback_meses", float("inf"))
    roi_van = _coerce_float(roi_simulation.get("van_12_meses_soles"))
    roi_pct = _coerce_float(roi_simulation.get("roi_porcentaje"))

    total_alertas = (
        int(round(_coerce_float(anomaly_frame["alertas_totales"].sum())))
        if "alertas_totales" in anomaly_frame.columns
        else 0
    )
    total_fugas = (
        int(round(_coerce_float(anomaly_frame["fugas_fisicas"].sum())))
        if "fugas_fisicas" in anomaly_frame.columns
        else 0
    )
    total_fraudes = (
        int(round(_coerce_float(anomaly_frame["fraudes_comerciales"].sum())))
        if "fraudes_comerciales" in anomaly_frame.columns
        else 0
    )

    document = Document()
    _configure_styles(document)
    _configure_header_footer(document)
    document.core_properties.title = REPORT_TITLE
    document.core_properties.author = "Rene-Pongo"
    document.core_properties.subject = "Agua No Facturada - EPS Ilo S.A."
    document.core_properties.comments = (
        "Reporte generado por el SIOC - Sistema Integrado de Operaciones."
    )

    # ----------------------- Portada / encabezado -----------------------
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run(REPORT_TITLE)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = _NAVY
    title.paragraph_format.space_after = Pt(2)

    _add_paragraph(
        document,
        REPORT_SUBTITLE,
        size=10.5,
        color=_CYAN_DARK,
        bold=True,
        space_after=2,
    )
    _add_paragraph(
        document,
        (
            f"Emitido el {datetime.now():%d/%m/%Y a las %H:%M} · "
            f"{EPS_NAME} · Ilo, Moquegua, Perú"
        ),
        size=9,
        color=_GRAY,
        space_after=10,
    )

    # --------------------------- 1. Resumen -----------------------------
    _add_heading(document, "1. Resumen ejecutivo", level=1)
    _add_paragraph(
        document,
        (
            f"Durante el periodo analizado, {EPS_NAME} produjo "
            f"{_fmt_number(producido)} m³ y facturó "
            f"{_fmt_number(facturado)} m³. El Agua No Facturada (ANF) alcanzó "
            f"{_fmt_number(anf_total)} m³, equivalente al {_fmt_pct(anf_pct)} "
            f"del volumen producido, con una pérdida económica estimada de "
            f"{_fmt_money(perdida_economica)} y un costo evitable de energía y "
            f"químicos de {_fmt_money(costo_energia)}."
        ),
    )
    _add_paragraph(
        document,
        (
            f"El motor de detección identificó {_fmt_int(total_alertas)} alertas "
            f"operativas, entre ellas {_fmt_int(total_fugas)} eventos compatibles "
            f"con fugas físicas y {_fmt_int(total_fraudes)} con fraude comercial "
            "o subregistro. La valorización utiliza la tarifa media ponderada de "
            f"S/ {TARIFF_STRUCTURE['tarifa_media_ponderada']:.2f}/m³ y el costo "
            f"variable de producción de S/ "
            f"{TARIFF_STRUCTURE['costo_variable_produccion']:.2f}/m³."
        ),
    )

    kpi_rows = [
        ["Volumen producido", f"{_fmt_number(producido)} m³"],
        ["Volumen facturado", f"{_fmt_number(facturado)} m³"],
        ["Agua No Facturada (ANF)", f"{_fmt_number(anf_total)} m³"],
        ["Porcentaje de ANF", _fmt_pct(anf_pct)],
        ["Pérdida física estimada", f"{_fmt_number(perdida_fisica)} m³"],
        ["Pérdida comercial estimada", f"{_fmt_number(perdida_comercial)} m³"],
        ["Pérdida económica total", _fmt_money(perdida_economica)],
        ["Costo energía/químicos perdido", _fmt_money(costo_energia)],
        ["Semáforo normativo SUNASS", semaforo],
    ]
    _add_table(
        document,
        ["Indicador", "Valor"],
        kpi_rows,
        widths_cm=[10.5, 6.5],
        font_size=9,
        align_right_from=1,
    )

    # --------------------------- 2. Semáforo ----------------------------
    _add_heading(document, "2. Semáforo normativo SUNASS", level=1)
    _add_semaforo_banner(document, semaforo, anf_pct)
    _add_paragraph(document, _semaforo_explanation(semaforo))
    _add_table(
        document,
        ["Nivel", "Rango de ANF", "Lectura operativa"],
        [
            [
                "VERDE",
                f"≤ {NORMATIVE_LIMITS['anf_alerta_amarilla_pct']:.1f}%",
                "Dentro de la meta referencial.",
            ],
            [
                "AMARILLO",
                (
                    f"> {NORMATIVE_LIMITS['anf_alerta_amarilla_pct']:.1f}% y "
                    f"≤ {NORMATIVE_LIMITS['anf_alerta_critica_pct']:.1f}%"
                ),
                "Alerta: plan de reducción intensivo.",
            ],
            [
                "ROJO",
                f"> {NORMATIVE_LIMITS['anf_alerta_critica_pct']:.1f}%",
                "Crítico: intervención inmediata prioritaria.",
            ],
        ],
        widths_cm=[2.6, 4.8, 9.6],
        font_size=8.8,
        align_right_from=None,
    )

    # --------------------- 3. Balance por sector ------------------------
    _add_heading(document, "3. Balance hídrico y valorización por sector", level=1)
    if sector_frame.empty:
        _add_paragraph(
            document,
            "No se recibió el desagregado sectorial para el periodo analizado.",
        )
    else:
        sector_rows = []
        for _, row in sector_frame.iterrows():
            sector_rows.append(
                [
                    str(row.get("sector_name", row.get("sector_id", "—"))),
                    _fmt_number(row.get("volumen_producido_m3")),
                    _fmt_number(row.get("volumen_facturado_m3")),
                    _fmt_number(row.get("anf_m3")),
                    _fmt_pct(row.get("anf_pct")),
                    _fmt_money(row.get("perdida_economica_soles")),
                    _fmt_int(row.get("alertas_fugas")),
                    str(row.get("semaforo_sunass", "—")),
                ]
            )
        sector_rows.append(
            [
                "TOTAL",
                _fmt_number(producido),
                _fmt_number(facturado),
                _fmt_number(anf_total),
                _fmt_pct(anf_pct),
                _fmt_money(perdida_economica),
                _fmt_int(total_fugas),
                semaforo,
            ]
        )
        _add_table(
            document,
            [
                "Sector",
                "Producido (m³)",
                "Facturado (m³)",
                "ANF (m³)",
                "ANF %",
                "Pérdida (S/)",
                "Fugas",
                "Semáforo",
            ],
            sector_rows,
            widths_cm=[3.7, 2.2, 2.2, 2.0, 1.4, 2.4, 1.2, 1.9],
            font_size=7.6,
            align_right_from=1,
            total_row=True,
        )

        top_sector = sector_frame.sort_values(
            "perdida_economica_soles", ascending=False
        ).iloc[0]
        _add_paragraph(
            document,
            (
                f"El sector con mayor impacto económico es "
                f"{top_sector.get('sector_name', top_sector.get('sector_id', '—'))}, "
                f"con {_fmt_money(top_sector.get('perdida_economica_soles'))} de "
                "pérdida valorizada. La priorización de intervenciones debe "
                "considerar este orden, complementado con la criticidad de las "
                "alertas operativas."
            ),
            size=9.5,
        )

    # ------------------------ 4. Alertas IA -----------------------------
    _add_heading(document, "4. Alertas operativas y comerciales", level=1)
    if anomaly_frame.empty:
        _add_paragraph(
            document,
            "No se registraron alertas de anomalías en el periodo analizado.",
        )
    else:
        anomaly_rows = []
        for _, row in anomaly_frame.iterrows():
            anomaly_rows.append(
                [
                    str(row.get("sector_name", row.get("sector_id", "—"))),
                    _fmt_int(row.get("registros")),
                    _fmt_int(row.get("alertas_totales")),
                    _fmt_pct(row.get("alertas_porcentaje")),
                    _fmt_int(row.get("fugas_fisicas")),
                    _fmt_int(row.get("fraudes_comerciales")),
                    _fmt_int(row.get("alertas_cloro")),
                    _fmt_int(row.get("alertas_presion")),
                    _fmt_pct(row.get("anf_promedio_pct")),
                    str(row.get("severidad_maxima", "—")).upper(),
                ]
            )
        _add_table(
            document,
            [
                "Sector",
                "Registros",
                "Alertas",
                "% Alertas",
                "Fugas",
                "Fraudes",
                "Cloro",
                "Presión",
                "ANF prom.",
                "Severidad",
            ],
            anomaly_rows,
            widths_cm=[3.4, 1.5, 1.4, 1.5, 1.2, 1.4, 1.2, 1.3, 1.5, 1.6],
            font_size=7.5,
            align_right_from=1,
        )
        _add_paragraph(
            document,
            (
                "Las alertas se generan combinando reglas normativas "
                "SUNASS/DIGESA con un modelo no supervisado Isolation Forest. "
                "La columna Severidad resume el nivel máximo detectado por "
                "sector: normal, baja, media, alta o crítica."
            ),
            size=9,
        )

    # ------------------------------ 5. ROI ------------------------------
    _add_heading(document, "5. Simulación de inversión y retorno (ROI)", level=1)
    _add_paragraph(
        document,
        (
            "La simulación evalúa un plan de reducción de ANF con una inversión "
            "inicial y un objetivo de recuperación de volumen. Los ahorros se "
            "valorizan a la tarifa media ponderada y el VAN se descuenta a una "
            "tasa anual referencial del 12%."
        ),
        size=9.8,
    )
    roi_rows = [
        ["Volumen recuperado por mes", f"{_fmt_number(roi_recuperado)} m³"],
        ["Ahorro mensual", _fmt_money(roi_ahorro_mensual)],
        ["Ahorro anual", _fmt_money(roi_ahorro_anual)],
        ["Periodo de recuperación (payback)", _fmt_payback(roi_payback)],
        ["VAN a 12 meses", _fmt_money(roi_van)],
        ["ROI porcentual", _fmt_pct(roi_pct)],
    ]
    _add_table(
        document,
        ["Indicador de retorno", "Resultado"],
        roi_rows,
        widths_cm=[10.5, 6.5],
        font_size=9,
        align_right_from=1,
    )

    if roi_van > 0 and roi_pct > 0:
        verdict = (
            "Veredicto: el plan es económicamente viable. El VAN a 12 meses es "
            f"positivo ({_fmt_money(roi_van)}) y el ROI alcanza "
            f"{_fmt_pct(roi_pct)}. Se recomienda aprobar la inversión y "
            "monitorear el volumen efectivamente recuperado."
        )
    else:
        verdict = (
            "Veredicto: con los parámetros actuales el plan no recupera la "
            "inversión en el horizonte de 12 meses. Se recomienda ampliar el "
            "objetivo de reducción, reducir el costo de intervención o "
            "priorizar acciones de bajo costo y rápido retorno."
        )
    _add_paragraph(document, verdict, size=9.8, bold=True, color=_NAVY)

    # ---------------- 6. Conclusiones y recomendaciones -----------------
    _add_heading(document, "6. Conclusiones y recomendaciones", level=1)
    conclusions = [
        (
            f"El ANF del periodo asciende a {_fmt_number(anf_total)} m³ "
            f"({_fmt_pct(anf_pct)}), con un impacto económico de "
            f"{_fmt_money(perdida_economica)}."
        ),
        (
            f"El semáforo SUNASS se encuentra en nivel {semaforo}. "
            + _semaforo_explanation(semaforo)
        ),
        (
            f"Se detectaron {_fmt_int(total_fugas)} eventos de fuga física y "
            f"{_fmt_int(total_fraudes)} de fraude comercial, que sustentan la "
            "priorización de inspecciones nocturnas y auditoría de medidores."
        ),
        (
            "La inversión simulada presenta un payback de "
            f"{_fmt_payback(roi_payback)} y un ROI de {_fmt_pct(roi_pct)}."
        ),
    ]
    for conclusion in conclusions:
        _add_bullet(document, conclusion)

    _add_heading(document, "Recomendaciones priorizadas", level=2)
    recommendations = [
        (
            "Intervenir primero el sector con mayor pérdida económica, "
            "combinando sectorización hidráulica y búsqueda activa de fugas "
            "nocturnas con correlación acústica."
        ),
        (
            "Auditar la macromedición y el parque de micromedidores para "
            "reducir el subregistro comercial, especialmente en redes antiguas."
        ),
        (
            "Mantener el cloro residual libre por encima de "
            f"{NORMATIVE_LIMITS['cloro_min_mg_l']:.1f} mg/L y la turbiedad por "
            f"debajo de {NORMATIVE_LIMITS['turbiedad_max_ntu']:.1f} NTU en todos "
            "los puntos de control, con énfasis en la cola de red de Pacocha."
        ),
        (
            "Revisar la presión de servicio en cota alta de Pampa Inalámbrica "
            f"para sostener el rango normativo de "
            f"{NORMATIVE_LIMITS['presion_min_mca']:.0f} a "
            f"{NORMATIVE_LIMITS['presion_max_mca']:.0f} m.c.a."
        ),
        (
            "Institucionalizar el tablero SIOC como herramienta de seguimiento "
            "mensual del plan de reducción de ANF y de rendición de cuentas "
            "ante la gerencia y el regulador."
        ),
    ]
    for recommendation in recommendations:
        _add_bullet(document, recommendation)

    _add_paragraph(
        document,
        (
            "Nota metodológica: los datos operativos provienen del motor de "
            "telemetría del SIOC y de los modelos de detección y valorización "
            "de las fases previas. El reporte es reproducible y puede "
            "regenerarse con la semilla y el periodo seleccionados en el "
            "tablero."
        ),
        size=8,
        color=_GRAY,
        space_after=0,
    )

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
